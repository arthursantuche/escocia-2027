"""Gera "Portal Escócia 2027 - Dados.xlsx" a partir de abas.json (ver exportar-abas.js).

O arquivo é enviado ao Google Drive com conversão para Google Sheets. Mesmo conteúdo,
estilo e regras de apps-script/Setup.gs (inclusive as ARRAYFORMULA, que o Sheets
reconhece na conversão). A cor da fonte do corpo é aplicada só às linhas com dados.

Uso: python scripts/gerar-planilha.py abas.json saida.xlsx   (requer openpyxl)
"""
import json
import re
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# Tokens da Identidade Visual iPhone Duo.
INK = '1D1D1F'
STUDIO_MIST = 'F5F5F7'
HAIRLINE = 'D6D6D6'
APPLE_BLUE = '0066CC'
SLATE = '707070'
LAUNCH_ORANGE = 'B64400'
FONTE = 'Inter'

ULTIMA_LINHA = 1000      # alcance das validações


def valor(v):
    if isinstance(v, dict) and '$data' in v:
        return date.fromisoformat(v['$data'])
    return v


def montar(ws, aba, wb):
    n_col = len(aba['colunas'])
    cab_fonte = Font(name=FONTE, size=10, bold=True, color=INK)
    cab_fundo = PatternFill('solid', fgColor=STUDIO_MIST)
    cab_borda = Border(bottom=Side(style='thin', color=HAIRLINE))
    for c, nome in enumerate(aba['colunas'], start=1):
        cel = ws.cell(row=1, column=c, value=nome)
        cel.font, cel.fill, cel.border = cab_fonte, cab_fundo, cab_borda
        cel.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'A2'

    corpo = Font(name=FONTE, size=10, color=INK)
    link = Font(name=FONTE, size=10, color=APPLE_BLUE)
    colunas_link = {i + 1 for i, n in enumerate(aba['colunas']) if n.startswith('Link')}
    texto = set(aba.get('texto', []))
    for r, linha in enumerate(aba['linhas'], start=2):
        for c, v in enumerate(linha, start=1):
            cel = ws.cell(row=r, column=c, value=valor(v) if v != '' else None)
            cel.font = link if c in colunas_link else corpo
            cel.alignment = Alignment(vertical='top', wrap_text=True)
            if c in texto:
                cel.number_format = '@'
            elif isinstance(cel.value, date) and aba.get('datasNaColunaB'):
                cel.number_format = 'dd/mm/yyyy'

    for cel, formula in (aba.get('formulas') or {}).items():
        ws[cel] = formula
        ws[cel].font = corpo

    for col, fmt in (aba.get('formatos') or {}).items():
        if not fmt:
            continue
        fim = max(ws.max_row, 2)
        for r in range(2, fim + 1):
            ws[f'{col}{r}'].number_format = fmt

    for i, px in enumerate(aba.get('larguras', []), start=1):
        ws.column_dimensions[get_column_letter(i)].width = round(px / 7, 1)

    for col, opcoes in (aba.get('validacoes') or {}).items():
        dv = DataValidation(type='list', formula1='"' + ','.join(opcoes) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle='Valor inválido', error='Escolha uma opção da lista.')
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{ULTIMA_LINHA}')
    for col, intervalo in (aba.get('validacaoIntervalo') or {}).items():
        folha, faixa = intervalo.split('!')
        m = re.match(r'([A-Z]+)(\d+):([A-Z]+)', faixa)
        dv = DataValidation(type='list', formula1=f'{folha}!${m[1]}${m[2]}:${m[3]}${ULTIMA_LINHA}', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{ULTIMA_LINHA}')

    if aba['nome'] == 'Checklist':
        faixa = f'F2:F{ULTIMA_LINHA}'
        ws.conditional_formatting.add(faixa, CellIsRule(operator='equal', formula=['"⚠️ pendente"'], font=Font(color=LAUNCH_ORANGE)))
        ws.conditional_formatting.add(faixa, CellIsRule(operator='equal', formula=['"feito"'], font=Font(color=SLATE)))
    if aba['nome'] == 'Orcamento':
        for r in range(2, ws.max_row + 1):
            for c in 'CDE':
                ws[f'{c}{r}'].alignment = Alignment(horizontal='right', vertical='top')


def main(entrada, saida):
    abas = json.load(open(entrada, encoding='utf-8'))
    wb = Workbook()
    wb.remove(wb.active)
    for aba in abas:
        montar(wb.create_sheet(aba['nome']), aba, wb)

    viagem = next(a for a in abas if a['nome'] == 'Viagem')
    linha = next(i for i, l in enumerate(viagem['linhas'], start=2) if l[0] == 'cotacao_gbp_brl')
    wb.defined_names['cotacao'] = DefinedName('cotacao', attr_text=f'Viagem!$B${linha}')
    wb.save(saida)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
